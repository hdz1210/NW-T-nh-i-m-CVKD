import openpyxl, re, sys, datetime, math
sys.stdout.reconfigure(encoding='utf-8')

def parse_number_safe(val):
    if val is None or val == '': return 0.0
    if isinstance(val, (int, float)): return float(val)
    s = str(val).strip()
    if not s: return 0.0
    if '.' in s and ',' in s: s = s.replace('.', '').replace(',', '.')
    elif '.' in s and re.search(r'\.\d{3}', s): s = s.replace('.', '')
    elif ',' in s: s = s.replace(',', '.')
    try: return float(re.sub(r'[^\d.-]', '', s))
    except: return 0.0

wb = openpyxl.load_workbook('Copy of NW_RawData (1).xlsx', data_only=True)
ws_th = wb['Tổng hợp']
ws_data = wb['Data']

th_months = []
for c in range(10, ws_th.max_column + 1):
    val = ws_th.cell(row=3, column=c).value
    if isinstance(val, datetime.datetime): th_months.append({'col': c, 'key': val.strftime('%Y-%m')})
    elif isinstance(val, str) and '/' in val:
        p = val.split('/')
        if len(p) == 2: th_months.append({'col': c, 'key': f"{p[1]}-{p[0].zfill(2)}"})

all_th_rules = []
cur_cdt = ''
cur_code = ''
cur_name = ''
cur_region = ''
for r in range(4, ws_th.max_row + 1):
    cdt = ws_th.cell(row=r, column=2).value
    code = ws_th.cell(row=r, column=3).value
    name = ws_th.cell(row=r, column=4).value
    region = ws_th.cell(row=r, column=5).value
    sp = ws_th.cell(row=r, column=7).value
    lc = ws_th.cell(row=r, column=8).value
    kg = ws_th.cell(row=r, column=9).value

    if cdt: cur_cdt = str(cdt).strip()
    if code: cur_code = str(code).strip()
    if name: cur_name = str(name).strip()
    if region: cur_region = str(region).strip()

    if not cur_code and not cur_cdt: continue

    month_scores = {}
    for m in th_months:
        score_val = ws_th.cell(row=r, column=m['col']).value
        if score_val is not None and score_val != '':
            try: month_scores[m['key']] = float(score_val)
            except: pass

    all_th_rules.append({
        'row': r,
        'cdt': cur_cdt,
        'code': cur_code,
        'name': cur_name,
        'region': cur_region,
        'fund': 'Quỹ NW',
        'sanPham': str(sp or 'Tất cả').strip(),
        'loaiCan': str(lc or 'Tất cả').strip(),
        'khoangGia': str(kg or 'Tất cả').strip(),
        'monthScores': month_scores
    })

# Add Month 9 Rules
m9_rules = [
    # Masterise
    {'cdt': 'Masterise Homes', 'code': 'Tất cả', 'name': 'Tất cả', 'region': 'Tất cả', 'fund': 'Quỹ NW', 'sanPham': 'Tất cả', 'loaiCan': 'Tất cả', 'khoangGia': '< 6 tỷ', 'monthScores': {'2026-09': 3.0}},
    {'cdt': 'Masterise Homes', 'code': 'Tất cả', 'name': 'Tất cả', 'region': 'Tất cả', 'fund': 'Quỹ NW', 'sanPham': 'Tất cả', 'loaiCan': 'Tất cả', 'khoangGia': '6 - 10 tỷ', 'monthScores': {'2026-09': 4.0}},
    {'cdt': 'Masterise Homes', 'code': 'Tất cả', 'name': 'Tất cả', 'region': 'Tất cả', 'fund': 'Quỹ NW', 'sanPham': 'Tất cả', 'loaiCan': 'Tất cả', 'khoangGia': '10 - 20 tỷ', 'monthScores': {'2026-09': 5.0}},
    {'cdt': 'Masterise Homes', 'code': 'Tất cả', 'name': 'Tất cả', 'region': 'Tất cả', 'fund': 'Quỹ NW', 'sanPham': 'Tất cả', 'loaiCan': 'Tất cả', 'khoangGia': '20 - 30 tỷ', 'monthScores': {'2026-09': 6.0}},
    {'cdt': 'Masterise Homes', 'code': 'Tất cả', 'name': 'Tất cả', 'region': 'Tất cả', 'fund': 'Quỹ NW', 'sanPham': 'Tất cả', 'loaiCan': 'Tất cả', 'khoangGia': '30 - 50 tỷ', 'monthScores': {'2026-09': 7.0}},
    {'cdt': 'Masterise Homes', 'code': 'Tất cả', 'name': 'Tất cả', 'region': 'Tất cả', 'fund': 'Quỹ NW', 'sanPham': 'Tất cả', 'loaiCan': 'Tất cả', 'khoangGia': '> 50 tỷ', 'monthScores': {'2026-09': 7.0}},
    # CĐT Khác
    {'cdt': 'Các CĐT Khác', 'code': 'Tất cả', 'name': 'Tất cả', 'region': 'Tất cả', 'fund': 'Quỹ NW', 'sanPham': 'Tất cả', 'loaiCan': 'Tất cả', 'khoangGia': '< 10 tỷ', 'monthScores': {'2026-09': 2.0}},
    {'cdt': 'Các CĐT Khác', 'code': 'Tất cả', 'name': 'Tất cả', 'region': 'Tất cả', 'fund': 'Quỹ NW', 'sanPham': 'Tất cả', 'loaiCan': 'Tất cả', 'khoangGia': '10 - 20 tỷ', 'monthScores': {'2026-09': 3.0}},
    {'cdt': 'Các CĐT Khác', 'code': 'Tất cả', 'name': 'Tất cả', 'region': 'Tất cả', 'fund': 'Quỹ NW', 'sanPham': 'Tất cả', 'loaiCan': 'Tất cả', 'khoangGia': '20 - 30 tỷ', 'monthScores': {'2026-09': 4.0}},
    {'cdt': 'Các CĐT Khác', 'code': 'Tất cả', 'name': 'Tất cả', 'region': 'Tất cả', 'fund': 'Quỹ NW', 'sanPham': 'Tất cả', 'loaiCan': 'Tất cả', 'khoangGia': '30 - 50 tỷ', 'monthScores': {'2026-09': 5.0}},
    {'cdt': 'Các CĐT Khác', 'code': 'Tất cả', 'name': 'Tất cả', 'region': 'Tất cả', 'fund': 'Quỹ NW', 'sanPham': 'Tất cả', 'loaiCan': 'Tất cả', 'khoangGia': '> 50 tỷ', 'monthScores': {'2026-09': 5.0}},
    # Quỹ Chéo
    {'cdt': 'Tất cả', 'code': 'Tất cả', 'name': 'Tất cả', 'region': 'Tất cả', 'fund': 'Quỹ Chéo', 'sanPham': 'Tất cả', 'loaiCan': 'Tất cả', 'khoangGia': '< 20 tỷ', 'monthScores': {'2026-09': 1.0}},
    {'cdt': 'Tất cả', 'code': 'Tất cả', 'name': 'Tất cả', 'region': 'Tất cả', 'fund': 'Quỹ Chéo', 'sanPham': 'Tất cả', 'loaiCan': 'Tất cả', 'khoangGia': '20 - 50 tỷ', 'monthScores': {'2026-09': 2.0}},
    {'cdt': 'Tất cả', 'code': 'Tất cả', 'name': 'Tất cả', 'region': 'Tất cả', 'fund': 'Quỹ Chéo', 'sanPham': 'Tất cả', 'loaiCan': 'Tất cả', 'khoangGia': '> 50 tỷ', 'monthScores': {'2026-09': 2.0}},
]
all_th_rules.extend(m9_rules)

ws_f2 = wb['Dự án F2']
f2_months = []
for c in range(3, ws_f2.max_column + 1):
    val = ws_f2.cell(row=3, column=c).value
    if isinstance(val, datetime.datetime): f2_months.append({'col': c, 'key': val.strftime('%Y-%m')})
    elif isinstance(val, str) and '/' in val:
        p = val.split('/')
        if len(p) == 2: f2_months.append({'col': c, 'key': f"{p[1]}-{p[0].zfill(2)}"})

f2_map = {}
for r in range(4, ws_f2.max_row + 1):
    name = ws_f2.cell(row=r, column=1).value
    if not name: continue
    m_scores = {}
    for m in f2_months:
        s_val = ws_f2.cell(row=r, column=m['col']).value
        if s_val is not None and s_val != '':
            try: m_scores[m['key']] = float(s_val)
            except: pass
    f2_map[str(name).strip().lower()] = m_scores

mas_vcg = set()
if 'Danh sách căn MAS VCG' in wb.sheetnames:
    ws_mas = wb['Danh sách căn MAS VCG']
    for r in range(2, ws_mas.max_row + 1):
        v1, v3 = ws_mas.cell(row=r, column=1).value, ws_mas.cell(row=r, column=3).value
        if v1: mas_vcg.add(str(v1).strip().upper())
        if v3: mas_vcg.add(str(v3).strip().upper())

gian_xay = {}
if 'Giãn xây HVB' in wb.sheetnames:
    ws_gx = wb['Giãn xây HVB']
    for r in range(2, ws_gx.max_row + 1):
        v1, v2 = ws_gx.cell(row=r, column=1).value, ws_gx.cell(row=r, column=2).value
        if v1 and v2 is not None:
            try: gian_xay[str(v1).strip().upper()] = float(v2)
            except: pass

cbnv = {}
if 'CBNV' in wb.sheetnames:
    ws_cb = wb['CBNV']
    for r in range(2, ws_cb.max_row + 1):
        v2, v3 = ws_cb.cell(row=r, column=2).value, ws_cb.cell(row=r, column=3).value
        if v2 and v3: cbnv[str(v2).strip().upper()] = str(v3).strip().upper()

def normalize_unit(u):
    if not u: return ''
    s = str(u).strip().lower()
    s = re.sub(r'[\(\[\{].*?[\)\]\}]', '', s)
    s = s.replace('phòng ngủ', 'pn').replace('căn hộ', '').replace('can ho', '')
    return re.sub(r'\s+', ' ', s).strip()

def is_all(val):
    if not val: return True
    s = str(val).strip().lower()
    return s in ('', '*', 'tất cả', 'tat ca', 'all')

def match_rules_py(cand_sp, cand_lc, cand_gia, tx_sp, tx_lc, tx_price):
    tsp = (tx_sp or '').strip().lower()
    tlc = normalize_unit(tx_lc)
    p = tx_price or 0.0

    if not is_all(cand_sp):
        sp = cand_sp.strip().lower()
        if 'thấp tầng' in sp and 'cao tầng' in tsp: return False
        if 'cao tầng' in sp and 'thấp tầng' in tsp: return False

    if not is_all(cand_gia):
        gia = cand_gia.strip().lower()
        rm = re.search(r'(?:từ\s*)?(\d+(?:[.,]\d+)?)\s*(?:-|đến)\s*(\d+(?:[.,]\d+)?)', gia)
        if rm:
            min_p = float(rm.group(1).replace(',', '.'))
            max_p = float(rm.group(2).replace(',', '.'))
            if max_p == 50.0 and p > 50.0: pass
            elif p < min_p or p > max_p: return False
        else:
            fm = re.search(r'(?:>=|>|trên|từ)\s*(\d+(?:[.,]\d+)?)', gia)
            if fm:
                min_p = float(fm.group(1).replace(',', '.'))
                if p < min_p: return False
            else:
                tm = re.search(r'(?:<=|<|dưới)\s*(\d+(?:[.,]\d+)?)', gia)
                if tm:
                    max_p = float(tm.group(1).replace(',', '.'))
                    if p >= max_p: return False

    if not is_all(cand_lc):
        lc = cand_lc.strip().lower()
        tokens = [normalize_unit(t) for t in re.split(r'[,;\n]|(?:\s+và\s+)', lc) if t.strip()]
        if tokens:
            matched = False
            for tok in tokens:
                if tlc == tok: matched = True; break
                if tok in ('shophouse', 'shophouses') and ('shophouse' in tlc or tlc in ('shop đế', 'tmdv')):
                    matched = True; break
                if tok == '1pn' and tlc in ('1pn', '1pn+'): matched = True; break
                if tok == '2pn' and tlc == '2pn': matched = True; break
                if tok == '3pn' and (tlc == '3pn' or '3br' in tlc or '3pn+' in tlc): matched = True; break
                if tok == '4pn' and (tlc == '4pn' or '4br' in tlc): matched = True; break
            if not matched: return False

    return True

def parse_date_safe(val, row):
    if row and row[2] is not None and row[3] is not None:
        try:
            y = int(row[3])
            m = int(row[2])
            d = 1
            if row[1] is not None:
                pd = int(row[1])
                if 1 <= pd <= 31: d = pd
            return datetime.datetime(y, m, d)
        except: pass
    if isinstance(val, datetime.datetime): return val
    if isinstance(val, str):
        val = val.strip()
        m1 = re.match(r'^(\d{1,2})[\/\-](\d{1,2})[\/\-](\d{4})', val)
        if m1: return datetime.datetime(int(m1.group(3)), int(m1.group(2)), int(m1.group(1)))
        m2 = re.match(r'^(\d{4})[\/\-](\d{1,2})[\/\-](\d{1,2})', val)
        if m2: return datetime.datetime(int(m2.group(1)), int(m2.group(2)), int(m2.group(3)))
    return None

def match_cdt(rule_cdt, tx_cdt):
    if is_all(rule_cdt): return True
    rc = rule_cdt.strip().lower()
    tc = (tx_cdt or '').strip().lower()
    if not tc: return True
    if 'khác' in rc or 'khac' in rc:
        return 'masterise' not in tc
    tokens = [t.strip().lower() for t in re.split(r'[,;]', rule_cdt) if t.strip()]
    return any(t in tc or tc in t for t in tokens)

def match_region(rule_region, tx_region):
    if is_all(rule_region): return True
    rr = rule_region.strip().lower()
    tr = (tx_region or '').strip().lower()
    if not tr: return True
    tokens = [t.strip().lower() for t in re.split(r'[,;]', rule_region) if t.strip()]
    return any(t in tr or tr in t for t in tokens)

def match_project(rule_code, rule_name, du_an, ten_du_an):
    if is_all(rule_code): return True
    da = du_an.strip().lower()
    tokens = [t.strip().lower() for t in re.split(r'[,;]', rule_code) if t.strip()]
    if any(t == da or t in da or da in t for t in tokens):
        return True
    if not is_all(rule_name):
        rn = rule_name.strip().lower()
        tda = (ten_du_an or '').strip().lower()
        if rn and tda and (rn in tda or tda in rn):
            return True
    return False

def evaluate_row_hierarchical(row):
    date_bc = parse_date_safe(row[0], row)
    if not date_bc: return ''

    raw_loai_quy = str(row[25] or '').strip()
    if not raw_loai_quy: return ''
    loai_quy = 'Quỹ NW' if 'NW' in raw_loai_quy.upper() else 'Quỹ Chéo'

    du_an = str(row[5] or '').strip()
    ma_can = str(row[6] or '').strip().upper()
    pkd = str(row[7] or '').strip()
    cvkd_name = str(row[8] or '').strip().upper()
    trang_thai = str(row[9] or '').strip()
    gia_chua_vat = parse_number_safe(row[10])
    gia_gom_vat = parse_number_safe(row[11])
    san_pham = str(row[14] or '').strip()
    loai_can = str(row[16] or '').strip()
    ma_nv = str(row[17] or '').strip().upper()
    tx_cdt = str(row[22] or '').strip()
    tx_mien = str(row[4] or '').strip()
    ten_du_an = str(row[28] or '').strip()
    ghi_chu = str(row[32] or '').strip()

    if trang_thai == 'Hủy': return 0
    if pkd == 'CTV/ĐỐI TÁC' or re.search(r'BLĐ|BO', pkd, re.I): return 0

    month_key = date_bc.strftime('%Y-%m')
    raw_val = gia_gom_vat if gia_gom_vat > 0 else gia_chua_vat
    if du_an == 'VHHVB' and ma_can in gian_xay:
        raw_val = gian_xay[ma_can]
    val_in_billion = raw_val / 1e9

    base_score = 0

    # 1. QUY CHEO
    if loai_quy == 'Quỹ Chéo':
        du_an_key = du_an.lower()
        if du_an_key in f2_map and month_key in f2_map[du_an_key]:
            base_score = f2_map[du_an_key][month_key]
        elif du_an_key in f2_map and month_key < '2026-09':
            base_score = f2_map[du_an_key].get(f2_months[0]['key'], 1)
        else:
            # Check general Quỹ Chéo rules in all_th_rules matching month_key
            matched_qcheo = []
            for r in all_th_rules:
                if r.get('fund') == 'Quỹ Chéo' and month_key in r['monthScores']:
                    if match_rules_py(r['sanPham'], r['loaiCan'], r['khoangGia'], san_pham, loai_can, val_in_billion):
                        matched_qcheo.append(r)
            if matched_qcheo:
                r = matched_qcheo[0]
                base_score = r['monthScores'][month_key]
                if val_in_billion > 50:
                    extra = math.floor((val_in_billion - 50.0001) / 10) + 1
                    base_score += extra
            else:
                base_score = 1
    else:
        # 2. QUY NW
        # STEP A: Try to find a specific project match first (matching du_an exactly or substring)
        du_an_key = du_an.lower()
        project_rules = [r for r in all_th_rules if not is_all(r['code']) and (r['code'].lower() == du_an_key or du_an_key in r['code'].lower() or r['code'].lower() in du_an_key)]
        
        candidates = []
        if project_rules:
            for r in project_rules:
                if not match_rules_py(r['sanPham'], r['loaiCan'], r['khoangGia'], san_pham, loai_can, val_in_billion): continue
                
                # Check month score
                score = None
                if month_key in r['monthScores']:
                    score = r['monthScores'][month_key]
                elif month_key < '2026-09':
                    # Fallback for old months: nearest month <= month_key
                    for m in th_months:
                        if m['key'] <= month_key and m['key'] in r['monthScores']:
                            score = r['monthScores'][m['key']]
                            break
                    if score is None:
                        # Nearest any month
                        for m in th_months:
                            if m['key'] in r['monthScores']:
                                score = r['monthScores'][m['key']]
                                break

                if score is not None:
                    # Specificity
                    spec = 100
                    if not is_all(r['sanPham']): spec += 10
                    if not is_all(r['loaiCan']): spec += 10
                    if not is_all(r['khoangGia']): spec += 10
                    candidates.append((spec, r, score))

        # STEP B: If no specific project rule matched, or for general rules (code == 'Tất cả')
        if not candidates:
            # Look at general rules eligible for month_key
            general_rules = [r for r in all_th_rules if is_all(r['code']) and r.get('fund', 'Quỹ NW') != 'Quỹ Chéo']
            for r in general_rules:
                if month_key not in r['monthScores']: continue
                if not match_cdt(r['cdt'], tx_cdt): continue
                if not match_region(r['region'], tx_mien): continue
                if not match_rules_py(r['sanPham'], r['loaiCan'], r['khoangGia'], san_pham, loai_can, val_in_billion): continue

                spec = 0
                if not is_all(r['cdt']):
                    spec += 30 if ('khác' in r['cdt'].lower() or 'khac' in r['cdt'].lower()) else 50
                if not is_all(r['region']): spec += 20
                if not is_all(r['sanPham']): spec += 10
                if not is_all(r['loaiCan']): spec += 10
                if not is_all(r['khoangGia']): spec += 10
                candidates.append((spec, r, r['monthScores'][month_key]))

        if candidates:
            candidates.sort(key=lambda x: x[0], reverse=True)
            best_rule = candidates[0][1]
            base_score = candidates[0][2]
            # Progressive pricing for Month 9 if > 50 tỷ
            if month_key >= '2026-09' and val_in_billion > 50:
                if '> 50' in best_rule['khoangGia'] or '50' in best_rule['khoangGia']:
                    extra = math.floor((val_in_billion - 50.0001) / 10) + 1
                    base_score += extra
        else:
            base_score = 2

    # Bonus & multipliers
    bonus_score = 0
    if du_an.upper() == 'VCG' and ma_can in mas_vcg:
        base_score = 8

    cot_n_quy = str(row[13] or '').strip()
    if datetime.datetime(2025, 5, 1) <= date_bc < datetime.datetime(2026, 9, 1) and val_in_billion >= 30 and cot_n_quy != 'Check D' and du_an.upper() != 'VCG':
        bonus_score += 1

    base_val = base_score + bonus_score
    time_mult = 1
    if datetime.datetime(2026, 2, 14) <= date_bc <= datetime.datetime(2026, 2, 28):
        time_mult = 2

    ptdt_mult = 1
    if re.search(r'PTĐT', pkd, re.I):
        ver_name = cbnv.get(ma_nv)
        is_internal = re.search(r'Cơ chế nội bộ', ghi_chu, re.I)
        if (ver_name and ver_name == cvkd_name) or is_internal:
            ptdt_mult = 1
        else:
            ptdt_mult = 0.5

    return base_val * time_mult * ptdt_mult

print("Running verification against all rows in Data sheet...")
total = 0
diffs = 0
diff_examples = []
for r_idx, row in enumerate(ws_data.iter_rows(min_row=2, max_row=ws_data.max_row, values_only=True)):
    total += 1
    sheet_score = row[23]
    calc_score = evaluate_row_hierarchical(row)
    
    if sheet_score is None or sheet_score == '':
        if calc_score != '' and calc_score != 0:
            diffs += 1
            if len(diff_examples) < 10:
                diff_examples.append((r_idx + 2, row[0], row[5], row[6], sheet_score, calc_score))
    else:
        try:
            s_val = float(sheet_score)
            c_val = float(calc_score) if calc_score != '' else 0.0
            if abs(s_val - c_val) > 0.001:
                diffs += 1
                if len(diff_examples) < 10:
                    diff_examples.append((r_idx + 2, row[0], row[5], row[6], sheet_score, calc_score))
        except:
            if str(sheet_score) != str(calc_score):
                diffs += 1
                if len(diff_examples) < 10:
                    diff_examples.append((r_idx + 2, row[0], row[5], row[6], sheet_score, calc_score))

print(f"VERIFICATION RESULT: Tested {total} rows. Differences: {diffs}")
if diff_examples:
    print("Sample differences:")
    for ex in diff_examples:
        print(f"  Row {ex[0]}: Date={ex[1]}, DA={ex[2]}, Can={ex[3]}, Sheet={ex[4]}, Calc={ex[5]}")

print("\n--- TESTING MONTH 9 TRANSACTIONS ---")
test_cases_m9 = [
    {'desc': 'Masterise < 6 tỷ (4.5 tỷ)', 'row': ['02/09/2026', 2, 9, 2026, 'Miền Bắc', 'MAS NEW', 'A1', 'PKD 1', 'CVKD A', 'Đã duyệt', 4500000000, 4500000000, '', '', 'Cao tầng', '', '1PN', 'NV01', '', '', '', '', 'Masterise Homes', '', '', 'Quỹ NW', '', '', 'Dự án Mas Mới', '', '', '', ''], 'expected': 3.0},
    {'desc': 'Masterise 8 tỷ (6-10 tỷ)', 'row': ['02/09/2026', 2, 9, 2026, 'Miền Bắc', 'MAS NEW', 'A2', 'PKD 1', 'CVKD A', 'Đã duyệt', 8000000000, 8000000000, '', '', 'Cao tầng', '', '2PN', 'NV01', '', '', '', '', 'Masterise Homes', '', '', 'Quỹ NW', '', '', 'Dự án Mas Mới', '', '', '', ''], 'expected': 4.0},
    {'desc': 'Masterise 15 tỷ (10-20 tỷ)', 'row': ['02/09/2026', 2, 9, 2026, 'Miền Bắc', 'MAS NEW', 'A3', 'PKD 1', 'CVKD A', 'Đã duyệt', 15000000000, 15000000000, '', '', 'Cao tầng', '', '3PN', 'NV01', '', '', '', '', 'Masterise Homes', '', '', 'Quỹ NW', '', '', 'Dự án Mas Mới', '', '', '', ''], 'expected': 5.0},
    {'desc': 'Masterise 25 tỷ (20-30 tỷ)', 'row': ['02/09/2026', 2, 9, 2026, 'Miền Bắc', 'MAS NEW', 'A4', 'PKD 1', 'CVKD A', 'Đã duyệt', 25000000000, 25000000000, '', '', 'Cao tầng', '', 'Duplex', 'NV01', '', '', '', '', 'Masterise Homes', '', '', 'Quỹ NW', '', '', 'Dự án Mas Mới', '', '', '', ''], 'expected': 6.0},
    {'desc': 'Masterise 40 tỷ (30-50 tỷ)', 'row': ['02/09/2026', 2, 9, 2026, 'Miền Bắc', 'MAS NEW', 'A5', 'PKD 1', 'CVKD A', 'Đã duyệt', 40000000000, 40000000000, '', '', 'Cao tầng', '', 'Penthouse', 'NV01', '', '', '', '', 'Masterise Homes', '', '', 'Quỹ NW', '', '', 'Dự án Mas Mới', '', '', '', ''], 'expected': 7.0},
    {'desc': 'Masterise 55 tỷ (mỗi 10 tỷ +1)', 'row': ['02/09/2026', 2, 9, 2026, 'Miền Bắc', 'MAS NEW', 'A6', 'PKD 1', 'CVKD A', 'Đã duyệt', 55000000000, 55000000000, '', '', 'Thấp tầng', '', 'Biệt thự', 'NV01', '', '', '', '', 'Masterise Homes', '', '', 'Quỹ NW', '', '', 'Dự án Mas Mới', '', '', '', ''], 'expected': 8.0},
    {'desc': 'Masterise 75 tỷ (mỗi 10 tỷ +1)', 'row': ['02/09/2026', 2, 9, 2026, 'Miền Bắc', 'MAS NEW', 'A7', 'PKD 1', 'CVKD A', 'Đã duyệt', 75000000000, 75000000000, '', '', 'Thấp tầng', '', 'Biệt thự', 'NV01', '', '', '', '', 'Masterise Homes', '', '', 'Quỹ NW', '', '', 'Dự án Mas Mới', '', '', '', ''], 'expected': 10.0},
    {'desc': 'Vinhomes 6 tỷ (<10 tỷ)', 'row': ['02/09/2026', 2, 9, 2026, 'Miền Bắc', 'VIN NEW', 'B1', 'PKD 1', 'CVKD A', 'Đã duyệt', 6000000000, 6000000000, '', '', 'Cao tầng', '', '1PN', 'NV01', '', '', '', '', 'Vinhomes', '', '', 'Quỹ NW', '', '', 'Vinhomes Mới', '', '', '', ''], 'expected': 2.0},
    {'desc': 'Capitaland 12 tỷ (10-20 tỷ)', 'row': ['02/09/2026', 2, 9, 2026, 'Miền Nam', 'CL NEW', 'B2', 'PKD 1', 'CVKD A', 'Đã duyệt', 12000000000, 12000000000, '', '', 'Cao tầng', '', '2PN', 'NV01', '', '', '', '', 'Capitaland', '', '', 'Quỹ NW', '', '', 'The Senique Mới', '', '', '', ''], 'expected': 3.0},
    {'desc': 'Gamuda 58 tỷ (mỗi 10 tỷ +1)', 'row': ['02/09/2026', 2, 9, 2026, 'Miền Nam', 'GM NEW', 'B3', 'PKD 1', 'CVKD A', 'Đã duyệt', 58000000000, 58000000000, '', '', 'Thấp tầng', '', 'Biệt thự', 'NV01', '', '', '', '', 'Gamuda Land', '', '', 'Quỹ NW', '', '', 'Eaton Park Mới', '', '', '', ''], 'expected': 6.0},
    {'desc': 'Quỹ Chéo 15 tỷ (<20 tỷ)', 'row': ['02/09/2026', 2, 9, 2026, 'Miền Bắc', 'CHEO NEW', 'C1', 'PKD 1', 'CVKD A', 'Đã duyệt', 15000000000, 15000000000, '', '', 'Cao tầng', '', '2PN', 'NV01', '', '', '', '', 'Ecopark', '', '', 'Quỹ Chéo', '', '', 'Eco Park', '', '', '', ''], 'expected': 1.0},
    {'desc': 'Quỹ Chéo 35 tỷ (20-50 tỷ)', 'row': ['02/09/2026', 2, 9, 2026, 'Miền Bắc', 'CHEO NEW', 'C2', 'PKD 1', 'CVKD A', 'Đã duyệt', 35000000000, 35000000000, '', '', 'Thấp tầng', '', 'Liền kề', 'NV01', '', '', '', '', 'Ecopark', '', '', 'Quỹ Chéo', '', '', 'Eco Park', '', '', '', ''], 'expected': 2.0},
    {'desc': 'Quỹ Chéo 65 tỷ (>50 tỷ)', 'row': ['02/09/2026', 2, 9, 2026, 'Miền Bắc', 'CHEO NEW', 'C3', 'PKD 1', 'CVKD A', 'Đã duyệt', 65000000000, 65000000000, '', '', 'Thấp tầng', '', 'Biệt thự', 'NV01', '', '', '', '', 'Ecopark', '', '', 'Quỹ Chéo', '', '', 'Eco Park', '', '', '', ''], 'expected': 4.0},
]

m9_passed = 0
for tc in test_cases_m9:
    res = evaluate_row_hierarchical(tc['row'])
    if abs(float(res) - tc['expected']) < 0.001:
        print(f"  [PASS] {tc['desc']}: Expected={tc['expected']}, Result={res}")
        m9_passed += 1
    else:
        print(f"  [FAIL] {tc['desc']}: Expected={tc['expected']}, Result={res}")

print(f"\nMonth 9 Test Cases Passed: {m9_passed}/{len(test_cases_m9)}")
