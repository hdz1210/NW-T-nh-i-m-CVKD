import re
import os
import urllib.request
import base64
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'''<w:tcMar {nsdecls("w")}>
        <w:top w:w="{top}" w:type="dxa"/>
        <w:bottom w:w="{bottom}" w:type="dxa"/>
        <w:left w:w="{left}" w:type="dxa"/>
        <w:right w:w="{right}" w:type="dxa"/>
    </w:tcMar>''')
    tcPr.append(tcMar)

def set_table_borders(table, color="D1D5DB", sz="4", val="single"):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(f'''<w:tblBorders {nsdecls("w")}>
            <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:left w:val="none"/>
            <w:right w:val="none"/>
        </w:tblBorders>''')
        tblPr[0].append(borders)

def fetch_mermaid_png(code, cache_dir):
    os.makedirs(cache_dir, exist_ok=True)
    import hashlib
    h = hashlib.md5(code.encode('utf-8')).hexdigest()
    cache_path = os.path.join(cache_dir, f"mermaid_{h}.png")
    if os.path.exists(cache_path) and os.path.getsize(cache_path) > 500:
        return cache_path
    
    # Try mermaid.ink
    try:
        b64 = base64.urlsafe_b64encode(code.strip().encode('utf-8')).decode('ascii')
        url = f"https://mermaid.ink/img/{b64}"
        req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req, timeout=12) as res:
            data = res.read()
            if len(data) > 200:
                with open(cache_path, 'wb') as f:
                    f.write(data)
                return cache_path
    except Exception as e:
        print(f"Failed mermaid.ink: {e}")
    
    # Fallback to kroki
    try:
        import zlib
        compressed = zlib.compress(code.strip().encode('utf-8'), 9)
        b64kroki = base64.urlsafe_b64encode(compressed).decode('ascii')
        url = f"https://kroki.io/mermaid/png/{b64kroki}"
        req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req, timeout=12) as res:
            data = res.read()
            if len(data) > 200:
                with open(cache_path, 'wb') as f:
                    f.write(data)
                return cache_path
    except Exception as e:
        print(f"Failed kroki: {e}")
    
    return None

def add_formatted_text(paragraph, text, base_bold=False, base_color=None, base_italic=False, base_font="Segoe UI", base_size=10.5):
    # Split text by markdown links, bold, inline code, italics
    # Pattern captures: [link](url), **bold**, `code`, *italic*
    tokens = re.split(r'(\[[^\]]+\]\([^\)]+\)|\*\*.*?\*\*|\`.*?\`|\*.*?\*)', text)
    for token in tokens:
        if not token:
            continue
        
        # 1. Markdown Link: [label](url)
        link_match = re.match(r'^\[(.*?)\]\((.*?)\)$', token)
        if link_match:
            label, url = link_match.groups()
            run = paragraph.add_run()
            run.font.name = base_font
            run.font.size = Pt(base_size)
            run.text = label
            # If it is an internal anchor #..., format neatly
            if url.startswith('#'):
                run.bold = base_bold
                run.font.color.rgb = base_color if base_color else RGBColor(30, 58, 138)
            else:
                run.font.color.rgb = RGBColor(37, 99, 235) # blue-600
                run.underline = True
            continue

        run = paragraph.add_run()
        run.font.name = base_font
        run.font.size = Pt(base_size)
        if base_color:
            run.font.color.rgb = base_color
        run.bold = base_bold
        run.italic = base_italic
        
        if token.startswith('**') and token.endswith('**') and len(token) >= 4:
            run.text = token[2:-2]
            run.bold = True
        elif token.startswith('`') and token.endswith('`') and len(token) >= 2:
            run.text = token[1:-1]
            run.font.name = "Consolas"
            run.font.size = Pt(base_size * 0.92)
            run.font.color.rgb = RGBColor(194, 65, 12) # orange-700
        elif token.startswith('*') and token.endswith('*') and len(token) >= 2 and not token.startswith('**'):
            run.text = token[1:-1]
            run.italic = True
        else:
            run.text = token

def convert_md_to_docx(md_path, docx_path, img_dir, cache_dir):
    with open(md_path, 'r', encoding='utf-8') as f:
        content = f.read()

    doc = Document()
    
    # Page setup
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)

    lines = content.split('\n')
    i = 0
    in_code_block = False
    code_lang = ""
    code_lines = []
    in_toc_section = False

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # Handle Code blocks & Mermaid
        if stripped.startswith('```'):
            if not in_code_block:
                in_code_block = True
                code_lang = stripped[3:].strip().lower()
                code_lines = []
                i += 1
                continue
            else:
                in_code_block = False
                code_str = '\n'.join(code_lines)
                if code_lang == 'mermaid':
                    png_path = fetch_mermaid_png(code_str, cache_dir)
                    if png_path and os.path.exists(png_path):
                        p = doc.add_paragraph()
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        p.paragraph_format.space_before = Pt(8)
                        p.paragraph_format.space_after = Pt(8)
                        run = p.add_run()
                        run.add_picture(png_path, width=Inches(6.2))
                    else:
                        table = doc.add_table(rows=1, cols=1)
                        table.alignment = WD_TABLE_ALIGNMENT.CENTER
                        set_cell_background(table.cell(0, 0), "F3F4F6")
                        set_cell_margins(table.cell(0, 0), top=120, bottom=120, left=160, right=160)
                        set_table_borders(table, color="E5E7EB")
                        cp = table.cell(0, 0).paragraphs[0]
                        cp.paragraph_format.space_before = Pt(2)
                        cp.paragraph_format.space_after = Pt(2)
                        run = cp.add_run(code_str)
                        run.font.name = "Consolas"
                        run.font.size = Pt(8.5)
                        run.font.color.rgb = RGBColor(55, 65, 81)
                else:
                    table = doc.add_table(rows=1, cols=1)
                    table.alignment = WD_TABLE_ALIGNMENT.CENTER
                    set_cell_background(table.cell(0, 0), "F8FAFC")
                    set_cell_margins(table.cell(0, 0), top=120, bottom=120, left=160, right=160)
                    set_table_borders(table, color="CBD5E1")
                    cp = table.cell(0, 0).paragraphs[0]
                    cp.paragraph_format.space_before = Pt(2)
                    cp.paragraph_format.space_after = Pt(2)
                    run = cp.add_run(code_str)
                    run.font.name = "Consolas"
                    run.font.size = Pt(9.0)
                    run.font.color.rgb = RGBColor(30, 41, 59)
                i += 1
                continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        # Empty line
        if not stripped:
            i += 1
            continue

        # Horizontal rule
        if stripped == '---':
            in_toc_section = False
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            p_border = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="1" w:color="E2E8F0"/></w:pBdr>')
            p._element.get_or_add_pPr().append(p_border)
            i += 1
            continue

        # Headings
        if stripped.startswith('# '):
            in_toc_section = False
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(8)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(stripped[2:])
            run.font.name = "Segoe UI"
            run.font.size = Pt(20)
            run.bold = True
            run.font.color.rgb = RGBColor(30, 58, 138) # #1e3a8a
            i += 1
            continue
        elif stripped.startswith('## '):
            heading_title = stripped[3:].strip()
            if 'MỤC LỤC' in heading_title.upper():
                in_toc_section = True
            else:
                in_toc_section = False

            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(heading_title)
            run.font.name = "Segoe UI"
            run.font.size = Pt(14.5)
            run.bold = True
            run.font.color.rgb = RGBColor(29, 78, 216) # #1d4ed8
            i += 1
            continue
        elif stripped.startswith('### '):
            in_toc_section = False
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(stripped[4:])
            run.font.name = "Segoe UI"
            run.font.size = Pt(12)
            run.bold = True
            run.font.color.rgb = RGBColor(3, 105, 161) # #0369a1
            i += 1
            continue
        elif stripped.startswith('#### '):
            in_toc_section = False
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(stripped[5:])
            run.font.name = "Segoe UI"
            run.font.size = Pt(11)
            run.bold = True
            run.font.color.rgb = RGBColor(51, 65, 85) # #334155
            i += 1
            continue

        # Markdown Images: ![alt](url)
        img_match = re.match(r'^!\[(.*?)\]\((.*?)\)$', stripped)
        if img_match:
            alt_text, img_rel_path = img_match.groups()
            local_img_path = os.path.normpath(os.path.join(os.path.dirname(md_path), img_rel_path))
            if os.path.exists(local_img_path):
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(8)
                p.paragraph_format.space_after = Pt(4)
                run = p.add_run()
                run.add_picture(local_img_path, width=Inches(6.2))
            i += 1
            continue

        # Image Caption or Italic note: *caption*
        if stripped.startswith('*') and stripped.endswith('*') and not stripped.startswith('**'):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(8)
            caption_text = stripped[1:-1]
            add_formatted_text(p, caption_text, base_italic=True, base_color=RGBColor(100, 116, 139), base_size=9.5)
            i += 1
            continue

        # Blockquote: > text
        if stripped.startswith('> '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.left_indent = Inches(0.3)
            p_border = parse_xml(f'<w:pBdr {nsdecls("w")}><w:left w:val="single" w:sz="18" w:space="10" w:color="2563EB"/></w:pBdr>')
            p._element.get_or_add_pPr().append(p_border)
            add_formatted_text(p, stripped[2:], base_color=RGBColor(30, 41, 59), base_size=10.0)
            i += 1
            continue

        # Markdown Tables: | col1 | col2 |
        if stripped.startswith('|') and stripped.endswith('|'):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith('|') and lines[i].strip().endswith('|'):
                table_lines.append(lines[i].strip())
                i += 1
            
            if len(table_lines) >= 2:
                headers = [c.strip() for c in table_lines[0].strip('|').split('|')]
                data_rows = []
                for row_line in table_lines[2:]:
                    cols = [c.strip() for c in row_line.strip('|').split('|')]
                    data_rows.append(cols)

                table = doc.add_table(rows=len(data_rows) + 1, cols=len(headers))
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                set_table_borders(table, color="CBD5E1", sz="4")

                for c_idx, h_text in enumerate(headers):
                    cell = table.cell(0, c_idx)
                    set_cell_background(cell, "1E3A8A")
                    set_cell_margins(cell, top=140, bottom=140, left=140, right=140)
                    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                    cp = cell.paragraphs[0]
                    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    cp.paragraph_format.space_before = Pt(2)
                    cp.paragraph_format.space_after = Pt(2)
                    add_formatted_text(cp, h_text, base_bold=True, base_color=RGBColor(255, 255, 255), base_size=9.5)

                for r_idx, row_data in enumerate(data_rows):
                    bg_color = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
                    for c_idx in range(len(headers)):
                        val = row_data[c_idx] if c_idx < len(row_data) else ""
                        cell = table.cell(r_idx + 1, c_idx)
                        set_cell_background(cell, bg_color)
                        set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
                        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                        cp = cell.paragraphs[0]
                        cp.paragraph_format.space_before = Pt(2)
                        cp.paragraph_format.space_after = Pt(2)
                        add_formatted_text(cp, val, base_size=9.5, base_color=RGBColor(30, 41, 59))
                
                sp_p = doc.add_paragraph()
                sp_p.paragraph_format.space_before = Pt(0)
                sp_p.paragraph_format.space_after = Pt(4)
            continue

        # Bullet List & Numbered Items (including Table of Contents)
        bullet_match = re.match(r'^(\s*)([-*]|\d+\.)\s+(.*)$', line)
        if bullet_match:
            indent_str, marker, item_text = bullet_match.groups()
            
            # If item_text is a standalone markdown link: [Label](#anchor), clean it to Label
            link_clean = re.match(r'^\[(.*?)\]\([^\)]*\)$', item_text.strip())
            if link_clean:
                clean_label = link_clean.group(1)
            else:
                clean_label = item_text

            p = doc.add_paragraph()
            indent_level = len(indent_str) // 2
            
            if in_toc_section:
                # Custom clean styling for Table of Contents
                if marker.endswith('.'): # 1. 2. 3. ...
                    p.paragraph_format.space_before = Pt(3)
                    p.paragraph_format.space_after = Pt(2)
                    p.paragraph_format.left_indent = Inches(0.2)
                    num_run = p.add_run(f"{marker} ")
                    num_run.font.name = "Segoe UI"
                    num_run.font.size = Pt(11)
                    num_run.bold = True
                    num_run.font.color.rgb = RGBColor(29, 78, 216) # blue-700
                    add_formatted_text(p, clean_label, base_bold=True, base_size=11, base_color=RGBColor(30, 41, 59))
                else: # sub-items
                    p.paragraph_format.space_before = Pt(1)
                    p.paragraph_format.space_after = Pt(1)
                    p.paragraph_format.left_indent = Inches(0.45)
                    bullet_run = p.add_run('•  ')
                    bullet_run.font.name = "Arial"
                    bullet_run.font.size = Pt(9.5)
                    bullet_run.font.color.rgb = RGBColor(14, 165, 233) # sky-500
                    bullet_run.bold = True
                    add_formatted_text(p, clean_label, base_size=10, base_color=RGBColor(51, 65, 85))
            else:
                # General list items
                p.paragraph_format.space_before = Pt(1.5)
                p.paragraph_format.space_after = Pt(1.5)
                p.paragraph_format.left_indent = Inches(0.25 * (indent_level + 1))
                
                if marker in ['-', '*']:
                    bullet_run = p.add_run('•  ' if indent_level == 0 else '○  ')
                    bullet_run.font.name = "Arial"
                    bullet_run.font.size = Pt(10)
                    bullet_run.font.color.rgb = RGBColor(37, 99, 235)
                    bullet_run.bold = True
                else:
                    num_run = p.add_run(f"{marker} ")
                    num_run.font.name = "Segoe UI"
                    num_run.font.size = Pt(10)
                    num_run.bold = True
                    num_run.font.color.rgb = RGBColor(30, 58, 138)
                
                add_formatted_text(p, clean_label, base_size=10.0, base_color=RGBColor(30, 41, 59))
            
            i += 1
            continue

        # Standard Paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        add_formatted_text(p, stripped, base_size=10.5, base_color=RGBColor(30, 41, 59))
        i += 1

    doc.save(docx_path)
    print(f"Successfully exported to {docx_path}")

if __name__ == "__main__":
    md_file = r"c:\Users\Admin\Desktop\Diem_CVKD\HUONG_DAN_SU_DUNG.md"
    docx_file = r"c:\Users\Admin\Desktop\Diem_CVKD\HUONG_DAN_SU_DUNG.docx"
    img_folder = r"c:\Users\Admin\Desktop\Diem_CVKD\docs\images"
    cache_folder = r"c:\Users\Admin\Desktop\Diem_CVKD\docs\images\mermaid_cache"
    convert_md_to_docx(md_file, docx_file, img_folder, cache_folder)
