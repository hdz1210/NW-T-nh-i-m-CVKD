import openpyxl, re, sys
sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('Copy of NW_RawData (1).xlsx', data_only=True)
ws = wb['Tổng hợp']

def parse_price(gia_str):
    if not gia_str:
        return {'min': None, 'max': None, 'type': 'ALL', 'label': 'Tất cả', 'condition': 'Không giới hạn'}
    s = str(gia_str).strip()
    s_low = s.lower()
    if s_low in ['', '*', 'tất cả', 'tat ca', 'all', 'none']:
        return {'min': None, 'max': None, 'type': 'ALL', 'label': 'Tất cả', 'condition': 'Không giới hạn'}
    
    is_dat = 'giá đất' in s_low
    
    # 1. Range: Từ X-Y tỷ or X - Y tỷ
    m_range = re.search(r'(?:từ\s*)?(\d+(?:[.,]\d+)?)\s*(?:-|đến)\s*(\d+(?:[.,]\d+)?)\s*t[ỷỉ]', s_low)
    if m_range:
        min_p = float(m_range.group(1).replace(',', '.'))
        max_p = float(m_range.group(2).replace(',', '.'))
        return {
            'min': min_p, 
            'max': max_p, 
            'type': 'RANGE', 
            'is_dat': is_dat, 
            'condition': f'{min_p:g} tỷ <= Giá <= {max_p:g} tỷ', 
            'label': f'Từ {min_p:g}-{max_p:g} tỷ'
        }
    
    # 2. Above: Trên X tỷ or Từ X tỷ trở lên or >= X tỷ
    m_above = re.search(r'(?:từ\s*(\d+(?:[.,]\d+)?)\s*t[ỷỉ]\s*trở\s*lên)|(?:trên\s*(\d+(?:[.,]\d+)?)\s*t[ỷỉ])|(?:\>=\s*(\d+(?:[.,]\d+)?)\s*t[ỷỉ])', s_low)
    if m_above:
        min_p = float((m_above.group(1) or m_above.group(2) or m_above.group(3)).replace(',', '.'))
        prefix = 'Giá đất ' if is_dat else ''
        return {
            'min': min_p, 
            'max': None, 
            'type': 'ABOVE', 
            'is_dat': is_dat, 
            'condition': f'{prefix}Giá >= {min_p:g} tỷ', 
            'label': f'{prefix}Trên {min_p:g} tỷ'
        }
    
    # 3. Below: Dưới X tỷ or <= X tỷ or < X tỷ
    m_below = re.search(r'(?:dưới\s*(\d+(?:[.,]\d+)?)\s*t[ỷỉ])|(?:\<=\s*(\d+(?:[.,]\d+)?)\s*t[ỷỉ])|(?:\<\s*(\d+(?:[.,]\d+)?)\s*t[ỷỉ])', s_low)
    if m_below:
        max_p = float((m_below.group(1) or m_below.group(2) or m_below.group(3)).replace(',', '.'))
        prefix = 'Giá đất ' if is_dat else ''
        return {
            'min': None, 
            'max': max_p, 
            'type': 'BELOW', 
            'is_dat': is_dat, 
            'condition': f'{prefix}Giá < {max_p:g} tỷ', 
            'label': f'{prefix}Dưới {max_p:g} tỷ'
        }
        
    return {'min': None, 'max': None, 'type': 'UNKNOWN', 'condition': s, 'label': s}

cur_ma = ''
cur_da = ''
parsed_rows = []
for r in range(4, 110):
    c_ma = ws.cell(row=r, column=3).value
    c_da = ws.cell(row=r, column=4).value
    sp = ws.cell(row=r, column=7).value
    lc = ws.cell(row=r, column=8).value
    kg = ws.cell(row=r, column=9).value
    if c_ma: cur_ma = str(c_ma).strip()
    if c_da: cur_da = str(c_da).strip()
    
    raw_val = str(kg or 'Tất cả').strip()
    p = parse_price(raw_val)
    parsed_rows.append((r, cur_ma, cur_da, str(sp or 'Tất cả').strip(), str(lc or 'Tất cả').strip(), raw_val, p))

print(f'Total rows parsed: {len(parsed_rows)}')
print('\n--- ALL 11 ROWS WITH SPECIFIC PRICE RANGES ---')
for r, ma, da, sp, lc, raw_kg, p in parsed_rows:
    if p['type'] != 'ALL':
        print(f"Row {r:3d} | Mã: {ma:15s} | Raw: {raw_kg:20s} -> Parsed Condition: {p['condition']:25s} | Min: {str(p['min']):6s} | Max: {str(p['max']):6s}")

with open('scratch/parsed_106_rows.txt', 'w', encoding='utf-8') as f:
    f.write(f"{'STT':<4} | {'Dòng':<5} | {'Mã Dự Án':<16} | {'Khoảng Giá (Raw)':<22} | {'Khoảng Giá (Parse)':<22} | {'Toán Tử / Điều Kiện':<26} | {'Min (Tỷ)':<9} | {'Max (Tỷ)':<9}\n")
    f.write('-'*125 + '\n')
    for idx, (r, ma, da, sp, lc, raw_kg, p) in enumerate(parsed_rows, 1):
        min_str = f"{p['min']:g}" if p['min'] is not None else '-'
        max_str = f"{p['max']:g}" if p['max'] is not None else '-'
        f.write(f"{idx:<4} | {r:<5} | {ma:<16} | {raw_kg:<22} | {p['label']:<22} | {p['condition']:<26} | {min_str:<9} | {max_str:<9}\n")
